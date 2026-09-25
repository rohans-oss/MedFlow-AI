"""V9 — file readers for the upload connector: CSV (UTF-8, `,` `;` or tab) and Excel .xlsx (first sheet, or the sheet
named after the entity). Row 1 is the header; returned rows carry their spreadsheet line number for error reports."""

import csv
import hashlib
import io
from datetime import date, datetime

from app.integrations.engine import SyncError

MAX_ROWS = 50_000


def read(filename: str, content: bytes, entity: str) -> tuple[list[dict], list[int], str]:
    """→ (records, line numbers, sha256 of the file)."""
    digest = hashlib.sha256(content).hexdigest()
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows, lines = _xlsx(content, entity)
    elif name.endswith((".csv", ".txt")):
        rows, lines = _csv(content)
    elif name.endswith(".xls"):
        raise SyncError("Old .xls files are not supported — save the sheet as .xlsx or CSV")
    else:
        raise SyncError("Upload a .csv or .xlsx file")
    if not rows:
        raise SyncError("The file has no data rows")
    if len(rows) > MAX_ROWS:
        raise SyncError(f"At most {MAX_ROWS} rows per file")
    return rows, lines, digest


def _csv(content: bytes) -> tuple[list[dict], list[int]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise SyncError("The CSV file is not UTF-8 encoded") from e
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header: list[str] | None = None
    rows, lines = [], []
    for i, rec in enumerate(reader, start=1):
        if header is None:
            header = [h.strip() for h in rec]
            if not any(header):
                raise SyncError("The first line must be a header row")
            if len(set(header)) != len(header):
                raise SyncError("The header row has duplicate column names")
            continue
        if not any(c.strip() for c in rec):
            continue
        rows.append({h: (rec[j].strip() if j < len(rec) else "") for j, h in enumerate(header) if h})
        lines.append(i)
    return rows, lines


def _cell(v):
    if isinstance(v, datetime):
        return v.isoformat() if (v.hour, v.minute, v.second) != (0, 0, 0) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return v


def _xlsx(content: bytes, entity: str) -> tuple[list[dict], list[int]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — any parse failure is the user's file, not our bug
        raise SyncError("Could not read the Excel file (is it a valid .xlsx?)") from e
    ws = wb[entity] if entity in wb.sheetnames else wb.worksheets[0]
    header: list[str] | None = None
    rows, lines = [], []
    for i, rec in enumerate(ws.iter_rows(values_only=True), start=1):
        if header is None:
            header = [str(h).strip() if h is not None else "" for h in rec]
            if not any(header):
                raise SyncError("Row 1 of the sheet must be a header row")
            continue
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in rec):
            continue
        rows.append({h: _cell(rec[j]) if j < len(rec) else None for j, h in enumerate(header) if h})
        lines.append(i)
        if len(rows) > MAX_ROWS:
            break
    wb.close()
    return rows, lines
